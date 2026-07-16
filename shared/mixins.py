import io
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.contrib import messages
from django.shortcuts import redirect
from django.db.models import ProtectedError
from django.http import HttpResponse
from django.utils import timezone
from django.utils.html import escape
from django.db.models import Model as DjangoModel


class StaffRequiredMixin:
    """
    Mixin que verifica si el usuario es miembro del staff.
    Si no es staff, redirige con mensaje de error.
    """
    staff_redirect_url = '/'
    staff_error_message = 'You do not have permission to perform this action. Staff access required.'

    def dispatch(self, request, *args, **kwargs):
        # 1. El superusuario, staff o miembro del grupo Administrador siempre pasa
        if request.user.is_superuser or request.user.is_staff or request.user.groups.filter(name='Administrador').exists():
            return super().dispatch(request, *args, **kwargs)
        
        # 2. Si la vista define group_required y el usuario pertenece a alguno de ellos, pasa
        if hasattr(self, 'group_required') and request.user.groups.filter(name__in=self.group_required).exists():
            return super().dispatch(request, *args, **kwargs)

        messages.error(request, self.staff_error_message)
        return redirect(self.staff_redirect_url)


def check_export_permission(request, export_format):
    """
    Verifica si el usuario tiene el permiso global requerido para exportar
    (security.export_pdf / security.export_excel). Si no lo tiene, muestra un
    mensaje de error y retorna una redirección; si lo tiene, retorna None.

    Uso en vistas basadas en función (FBV) que exportan manualmente:
        redirect_response = check_export_permission(request, export_format)
        if redirect_response:
            return redirect_response
    """
    required_perm = 'security.export_pdf' if export_format == 'pdf' else 'security.export_excel'
    if not request.user.has_perm(required_perm):
        messages.error(request, 'No tiene permiso para exportar este reporte.')
        return redirect(request.path)
    return None


class ExportMixin:
    """
    Mixin genérico para vistas ListView que permite la exportación del QuerySet filtrado
    a formatos Excel y PDF.
    
    Se activa mediante los parámetros GET:
      ?export=excel
      ?export=pdf
      
    Las vistas secundarias pueden configurar opcionalmente:
      - export_fields: Lista de tuplas (nombre_campo, cabecera). Soporta relaciones con punto (ej. 'brand.name').
      - export_filename: Nombre base del archivo de salida.
    """
    export_fields = None
    export_filename = 'export'

    def get_export_fields(self):
        """Retorna los campos a exportar. Si no se especifican, se auto-detectan del modelo."""
        fields = self.export_fields
        if not fields:
            fields = [(f.name, str(f.verbose_name).capitalize()) for f in self.model._meta.fields]
            
        if hasattr(self, 'request') and self.request:
            cols = self.request.GET.getlist('columns')
            if cols:
                fields = [f for f in fields if f[0] in cols]
                
        return fields

    def get_export_filename(self):
        timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
        return f"{self.export_filename}_{timestamp}"

    def get_field_value(self, obj, field_path):
        """Obtiene el valor de un campo del objeto, soportando notación de punto y M2M."""
        parts = field_path.split('.')
        val = obj
        for part in parts:
            if val is None:
                break
            
            # Si es un manager de M2M
            if hasattr(val, 'all') and callable(getattr(val, 'all')):
                val = ", ".join(str(item) for item in val.all())
                break
                
            if hasattr(val, part):
                attr = getattr(val, part)
                if hasattr(attr, 'all') and callable(getattr(attr, 'all')):
                    val = attr
                elif callable(attr) and not isinstance(attr, DjangoModel):
                    try:
                        val = attr()
                    except Exception:
                        val = attr
                else:
                    val = attr
            else:
                return ""
        
        # Verificar si quedó un manager/queryset de M2M sin evaluar
        if hasattr(val, 'all') and callable(getattr(val, 'all')):
            val = ", ".join(str(item) for item in val.all())
            
        if val is None:
            return ""
        if isinstance(val, bool):
            return "Activo" if val else "Inactivo"
        return str(val)

    def get(self, request, *args, **kwargs):
        export_format = request.GET.get('export')
        if export_format in ['excel', 'pdf']:
            redirect_response = check_export_permission(request, export_format)
            if redirect_response:
                return redirect_response

            # Ejecutamos el queryset con todos los filtros actuales (sin paginar)
            self.object_list = self.get_queryset()
            fields = self.get_export_fields()
            filename = self.get_export_filename()
            
            if export_format == 'excel':
                return self.export_to_excel(self.object_list, fields, filename)
            elif export_format == 'pdf':
                return self.export_to_pdf(self.object_list, fields, filename)
                
        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        """Añade variables de control al contexto para facilitar los enlaces de exportación."""
        ctx = super().get_context_data(**kwargs)
        params = self.request.GET.copy()
        params.pop('page', None)
        ctx['query_string'] = params.urlencode()
        
        # Column selection feature
        fields = self.export_fields
        if not fields:
            fields = [(f.name, str(f.verbose_name).capitalize()) for f in self.model._meta.fields]
            
        ctx['available_columns'] = fields
        
        selected = self.request.GET.getlist('columns')
        if not selected:
            selected = [f[0] for f in fields]
        ctx['selected_columns'] = selected
        
        return ctx

    def export_to_excel(self, queryset, fields, filename):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Listado"
        
        # Encabezados
        headers = [f[1] if isinstance(f, tuple) else f for f in fields]
        ws.append(headers)
        
        # Datos
        field_paths = [f[0] if isinstance(f, tuple) else f for f in fields]
        for obj in queryset:
            row = []
            for path in field_paths:
                row.append(self.get_field_value(obj, path))
            ws.append(row)
            
        # Estilos en cabecera (Negrita y fondo oscuro premium)
        header_fill = PatternFill(start_color="343A40", end_color="343A40", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            
        # Auto-ajuste de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
            
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        wb.save(response)
        return response

    def export_to_pdf(self, queryset, fields, filename):
        buffer = io.BytesIO()
        
        headers = [f[1] if isinstance(f, tuple) else f for f in fields]
        field_paths = [f[0] if isinstance(f, tuple) else f for f in fields]
        num_cols = len(headers)
        
        # Si tiene más de 5 columnas, usamos orientación horizontal (Landscape)
        pagesize = landscape(letter) if num_cols > 5 else letter
        
        # Margen de 36pt (0.5 in)
        doc = SimpleDocTemplate(
            buffer,
            pagesize=pagesize,
            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
        )
        elements = []
        
        # Configurar estilos de ReportLab
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'PdfTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=16,
            textColor=colors.HexColor('#212529'),
            spaceAfter=15
        )
        normal_style = ParagraphStyle(
            'PdfBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=8,
            textColor=colors.HexColor('#495057')
        )
        header_style = ParagraphStyle(
            'PdfHeader',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=8,
            textColor=colors.white
        )
        
        # Agregar Título y Fecha
        plural_name = self.model._meta.verbose_name_plural.capitalize()
        elements.append(Paragraph(f"Reporte General de {plural_name}", title_style))
        elements.append(Paragraph(f"Generado el: {timezone.now().strftime('%d/%m/%Y %H:%M:%S')}", normal_style))
        elements.append(Spacer(1, 15))
        
        # Generar Datos de Tabla
        table_data = [[Paragraph(escape(h), header_style) for h in headers]]
        for obj in queryset:
            row = []
            for path in field_paths:
                val = self.get_field_value(obj, path)
                row.append(Paragraph(escape(val), normal_style))
            table_data.append(row)
            
        # Calcular ancho de columnas
        usable_width = pagesize[0] - 72
        col_width = usable_width / num_cols
        
        t = Table(table_data, colWidths=[col_width] * num_cols)
        
        # Estilos visuales de la tabla
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#343A40')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#DEE2E6')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor('#F8F9FA'), colors.white])
        ]))
        
        elements.append(t)
        doc.build(elements)
        
        pdf = buffer.getvalue()
        buffer.close()
        
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
        response.write(pdf)
        return response

class GroupRequiredMixin:
    """
    Mixin que verifica si el usuario pertenece a alguno
    de los roles (grupos) indicados en group_required.

    Uso:
        class GroupListView(LoginRequiredMixin, GroupRequiredMixin, ListView):
            group_required = ['Administrador']
    """
    group_required = []        # Lista de roles permitidos
    group_redirect_url = '/'   # A dónde redirigir si no tiene el rol
    group_error_message = 'No tiene permiso para acceder a esta opción.'

    def dispatch(self, request, *args, **kwargs):
        # 1. Si no inició sesión -> al login
        if not request.user.is_authenticated:
            return redirect('login')
        # 2. El superusuario siempre pasa
        if request.user.is_superuser:
            return super().dispatch(request, *args, **kwargs)
        # 3. ¿Pertenece a alguno de los roles permitidos?
        if request.user.groups.filter(name__in=self.group_required).exists():
            return super().dispatch(request, *args, **kwargs)
        # 4. No tiene el rol -> mensaje de error y redirección
        messages.error(request, self.group_error_message)
        return redirect(self.group_redirect_url)

class ProtectedDeleteMixin:
    """
    Mixin para DeleteView que atrapa ProtectedError (relaciones con
    on_delete=models.PROTECT) y muestra un mensaje amigable en vez de la
    página de error 500 cruda de Django.

    Nota técnica: en Django 4+/6, DeleteView llama a form_valid() en el POST,
    no a delete() -> por eso este mixin sobreescribe form_valid().
    """
    protected_error_message = "No se puede eliminar '{object}' porque está siendo utilizado en otros registros del sistema."

    def form_valid(self, form):
        self.object = self.get_object()
        object_repr = str(self.object)
        success_url = self.get_success_url()
        try:
            self.object.delete()
            messages.success(self.request, f"'{object_repr}' eliminado correctamente!")
        except ProtectedError:
            messages.error(self.request, self.protected_error_message.format(object=object_repr))
        return redirect(success_url)


class PermissionOrRedirectMixin:
    """
    Verificación de permisos finos (Django auth.Permission) para CBVs.
    A diferencia de PermissionRequiredMixin (que muestra un 403 crudo),
    este redirige con un mensaje de error, igual que GroupRequiredMixin,
    para mantener consistencia de UX en todo el sistema.

    Uso:
        class InvoiceDeleteView(PermissionOrRedirectMixin, LoginRequiredMixin, GroupRequiredMixin, DeleteView):
            permission_required = 'billing.delete_invoice'   # str o tupla de strings
    """
    permission_required = None
    permission_redirect_url = '/'
    permission_error_message = 'No tiene permiso para realizar esta acción.'

    def dispatch(self, request, *args, **kwargs):
        if not request.user.is_authenticated:
            return redirect('login')
        if request.user.is_superuser:
            return super().dispatch(request, *args, **kwargs)

        perms = self.permission_required
        if perms is None:
            return super().dispatch(request, *args, **kwargs)
        if isinstance(perms, str):
            perms = (perms,)

        if request.user.has_perms(perms):
            return super().dispatch(request, *args, **kwargs)

        messages.error(request, self.permission_error_message)
        return redirect(self.permission_redirect_url)
